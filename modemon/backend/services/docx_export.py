"""DOCX 导出：按排版规格生成中文 Word 论文。

三个中文 Word 的坑（都必须处理对，否则排版会明显不对）
----------------------------------------------------
1. **东亚字体必须单独设 `w:eastAsia`**：只设 `run.font.name` 只会作用于西文，
   中文仍然走默认字体。
2. **首行缩进要按「字符」而不是磅**：`w:ind/@w:firstLineChars` 以 1/100 字符为单位，
   这样才能随字号自适应；用磅写死会在改字号后错位。
3. **三线表要靠边框控制**：上/下粗、表头下细，中间无线——
   这是学术论文表格的规范形式，不能用 Word 内置样式凑。

规格字段全部取自 `typography`（恢复真值）。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

from . import paper_build as pb
from . import typography as ty
from .latex_export import markdown_body_to_tex  # 复用行内解析思路（此处仅取纯文本）

#: 1 磅 = 1/72 英寸 = 1.333 像素；python-docx 用 EMU，Cm/Pt 构造器已封装。


# ---------------------------------------------------------------------------
# 底层工具
# ---------------------------------------------------------------------------


def set_run_font(run, *, ascii_font: str, ea_font: str, size_pt: float,
                 bold: bool = False, color: Optional[str] = None) -> None:
    """设置 run 的字体，**同时设置西文与东亚字体**。"""
    run.font.name = ascii_font
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    # ★ 关键：东亚字体
    rfonts.set(qn("w:eastAsia"), ea_font)


def set_first_line_indent_chars(paragraph, chars: float) -> None:
    """首行缩进按**字符**设置（`w:firstLineChars`，单位 1/100 字符）。"""
    pf = paragraph.paragraph_format
    ppr = paragraph._p.get_or_add_pPr()
    ind = ppr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        ppr.append(ind)
    ind.set(qn("w:firstLineChars"), str(int(chars * 100)))
    # 同时给一个等效磅值兜底（部分阅读器不认 firstLineChars）
    pf.first_line_indent = Pt(chars * 12)  # 以 12pt 字号近似


def set_cell_borders(cell, *, top_pt: float = 0, bottom_pt: float = 0) -> None:
    """设置单元格上/下边框粗细（1/8 磅为单位）。"""
    tcpr = cell._tc.get_or_add_tcPr()
    borders = tcpr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcpr.append(borders)
    for tag, width_pt in (("w:top", top_pt), ("w:bottom", bottom_pt)):
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        if width_pt <= 0:
            el.set(qn("w:val"), "none")
        else:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(int(width_pt * 8)))   # sz 单位 1/8 磅
            el.set(qn("w:color"), "000000")


def shade_cell(cell, hex_color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


ALIGN = {
    "center": WD_ALIGN_PARAGRAPH.CENTER,
    "left": WD_ALIGN_PARAGRAPH.LEFT,
    "right": WD_ALIGN_PARAGRAPH.RIGHT,
    "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
}


# ---------------------------------------------------------------------------
# 文档构建
# ---------------------------------------------------------------------------


def _setup_page(doc: Document, p: ty.StyleProfile) -> None:
    for s in doc.sections:
        geo = p.page
        s.page_width = Cm(21.0)      # A4
        s.page_height = Cm(29.7)
        s.top_margin = Cm(float(geo["margin_top_cm"]))
        s.bottom_margin = Cm(float(geo["margin_bottom_cm"]))
        s.left_margin = Cm(float(geo["margin_left_cm"]))
        s.right_margin = Cm(float(geo["margin_right_cm"]))


def _normal_style(doc: Document, p: ty.StyleProfile) -> None:
    st = doc.styles["Normal"]
    body = p.body
    st.font.name = p.fonts["latin"]
    st.font.size = Pt(float(body["font_size_pt"]))
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:ascii"), p.fonts["latin"])
    rfonts.set(qn("w:hAnsi"), p.fonts["latin"])
    rfonts.set(qn("w:eastAsia"), p.fonts["chinese_body"])
    pf = st.paragraph_format
    pf.line_spacing = float(body["line_spacing"])
    pf.space_before = Pt(float(body["space_before_pt"]))
    pf.space_after = Pt(float(body["space_after_pt"]))


def add_paragraph(
    doc: Document, text: str, p: ty.StyleProfile, *,
    size_pt: Optional[float] = None, bold: bool = False,
    ea_font: Optional[str] = None, align: str = "left",
    indent_chars: float = 0.0, spacing: Optional[float] = None,
):
    para = doc.add_paragraph()
    para.alignment = ALIGN.get(align, WD_ALIGN_PARAGRAPH.LEFT)
    run = para.add_run(text)
    set_run_font(run, ascii_font=p.fonts["latin"],
                 ea_font=ea_font or p.fonts["chinese_body"],
                 size_pt=size_pt if size_pt is not None else float(p.body["font_size_pt"]),
                 bold=bold)
    if spacing is not None:
        para.paragraph_format.line_spacing = spacing
    if indent_chars:
        set_first_line_indent_chars(para, indent_chars)
    return para


def _strip_md(text: str) -> str:
    """去掉行内 Markdown 标记（docx 这里按纯文本写入）。"""
    import re
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"(?<!\*)\*([^*]+)\*(?!\*)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"^\s*[-*+]\s+", "· ", text)
    return text


def build_docx(
    doc_model: pb.PaperDoc,
    profile: ty.StyleProfile,
    *,
    include_tables: bool = True,
    include_figures: bool = True,
    include_ai_disclosure: bool = True,
) -> Document:
    """按规格构建 DOCX 文档对象。"""
    p = profile
    d = Document()
    _setup_page(d, p)
    _normal_style(d, p)

    # ---- 标题 ----
    t = p.title
    add_paragraph(d, doc_model.title, p,
                  size_pt=float(t["font_size_pt"]), bold=bool(t["bold"]),
                  ea_font=t["font_family"], align=str(t["alignment"]),
                  spacing=float(t["line_spacing"]))

    # ---- 摘要 ----
    if doc_model.abstract:
        ab = p.data.get("abstract", {})
        add_paragraph(d, "摘要", p, size_pt=float(ab.get("label_size_pt", 14)),
                      bold=bool(ab.get("label_bold", True)),
                      ea_font=p.fonts["chinese_heading"],
                      align=str(ab.get("label_alignment", "center")))
        add_paragraph(d, _strip_md(doc_model.abstract), p,
                      size_pt=float(ab.get("body_size_pt", 12)), indent_chars=2)

    # ---- 关键词 ----
    if doc_model.keywords:
        kw = p.data.get("keywords", {})
        para = d.add_paragraph()
        r1 = para.add_run("关键词")
        set_run_font(r1, ascii_font=p.fonts["latin"], ea_font=kw.get("font", "SimHei"),
                     size_pt=float(kw.get("label_size_pt", 12)),
                     bold=bool(kw.get("label_bold", True)))
        r2 = para.add_run("　" + "；".join(doc_model.keywords))
        set_run_font(r2, ascii_font=p.fonts["latin"], ea_font=p.fonts["chinese_body"],
                     size_pt=float(p.body["font_size_pt"]))

    # ---- 正文 ----
    for s in doc_model.sections:
        if s.level == 1 and doc_model.title and s.title.strip() == doc_model.title.strip() and not s.body.strip():
            continue
        add_paragraph(d, s.title, p, size_pt=p.heading_pt(s.level),
                      bold=bool(p.headings["bold"]),
                      ea_font=p.fonts["chinese_heading"],
                      align=str(p.headings.get(f"level{min(3, s.level)}_alignment", "left")))
        for line in s.body.splitlines():
            if not line.strip():
                continue
            add_paragraph(d, _strip_md(line), p,
                          indent_chars=0 if line.lstrip().startswith(("-", "*", "·"))
                          else float(p.body["first_line_indent_chars"]))

    # ---- 表格（程序生成，三线表） ----
    if include_tables:
        for tspec in doc_model.tables:
            add_paragraph(d, f"表 {tspec.name}　{tspec.caption}", p,
                          size_pt=float(p.table["font_size_pt"]) + 1,
                          bold=True, ea_font=p.fonts["chinese_heading"], align="center")
            tbl = d.add_table(rows=1 + len(tspec.rows), cols=len(tspec.headers))
            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
            tb = p.table
            hdr = tbl.rows[0].cells
            for i, h in enumerate(tspec.headers):
                hdr[i].text = ""
                para = hdr[i].paragraphs[0]
                para.alignment = ALIGN[str(tb["cell_alignment"])]
                run = para.add_run(str(h))
                set_run_font(run, ascii_font=p.fonts["latin"],
                             ea_font=p.fonts["chinese_heading"],
                             size_pt=float(tb["font_size_pt"]),
                             bold=bool(tb["header_bold"]))
                # 三线表：首行上粗线 + 表头下细线
                set_cell_borders(hdr[i], top_pt=float(tb["top_border_pt"]),
                                 bottom_pt=float(tb["header_border_pt"]))
            for r_i, row in enumerate(tspec.rows, start=1):
                cells = tbl.rows[r_i].cells
                for i, val in enumerate(row):
                    cells[i].text = ""
                    para = cells[i].paragraphs[0]
                    para.alignment = ALIGN[str(tb["cell_alignment"])]
                    run = para.add_run(str(val))
                    set_run_font(run, ascii_font=p.fonts["latin"],
                                 ea_font=p.fonts["chinese_body"],
                                 size_pt=float(tb["font_size_pt"]))
                    # 末行下粗线
                    if r_i == len(tspec.rows):
                        set_cell_borders(cells[i], bottom_pt=float(tb["bottom_border_pt"]))
            if tspec.source:
                add_paragraph(d, f"数据来源：{tspec.source}", p,
                              size_pt=float(tb["font_size_pt"]) - 1, align="center")

    # ---- 图 ----
    if include_figures:
        max_w = float(p.image["max_width_cm"])
        for f in doc_model.figures:
            path = Path(f.path)
            if path.is_file() and path.suffix.lower() in (".png", ".jpg", ".jpeg"):
                para = d.add_paragraph()
                para.alignment = ALIGN[str(p.image["alignment"])]
                try:
                    para.add_run().add_picture(str(path), width=Cm(max_w))
                except Exception:
                    pass
                cap = add_paragraph(d, f.display_caption, p,
                                    size_pt=float(p.body["font_size_pt"]) - 1,
                                    align="center")
                cap.paragraph_format.space_before = Pt(2)

    # ---- 参考文献 ----
    if doc_model.references:
        refs = p.references
        add_paragraph(d, "参考文献", p, size_pt=p.heading_pt(1), bold=True,
                      ea_font=p.fonts["chinese_heading"])
        for i, r in enumerate(doc_model.references, 1):
            para = d.add_paragraph()
            ind = OxmlElement("w:ind")
            ind.set(qn("w:left"), str(int(float(refs["hanging_indent_cm"]) * 567)))
            ind.set(qn("w:hanging"), str(int(float(refs["hanging_indent_cm"]) * 567)))
            para._p.get_or_add_pPr().append(ind)
            run = para.add_run(f"[{i}] {_strip_md(r)}")
            set_run_font(run, ascii_font=p.fonts["latin"], ea_font=p.fonts["chinese_body"],
                         size_pt=float(refs["font_size_pt"]))

    # ---- AI 使用申报 ----
    if include_ai_disclosure and doc_model.ai_disclosure:
        add_paragraph(d, "AI 工具使用情况", p, size_pt=p.heading_pt(1), bold=True,
                      ea_font=p.fonts["chinese_heading"])
        for i, rec in enumerate(doc_model.ai_disclosure, 1):
            line = (f"{i}. {rec.get('tool_name', '')}：{rec.get('purpose', '')}"
                    + (f"（采用情况：{rec.get('output_used')}）" if rec.get("output_used") else ""))
            add_paragraph(d, line, p, indent_chars=0)
    return d


# ---------------------------------------------------------------------------
# 端到端
# ---------------------------------------------------------------------------


def export_docx(
    workspace: Path,
    *,
    template: str = "comp_cumcm",
    title: str = "",
    abstract: str = "",
    keywords: Optional[list[str]] = None,
    out_name: str = "main.docx",
) -> dict[str, Any]:
    """从工作区装配并导出 DOCX。"""
    ws = Path(workspace)
    paper_dir = ws / "paper"
    paper_dir.mkdir(parents=True, exist_ok=True)

    doc_model, notes = pb.build_paper(ws, template=template, title=title,
                                      abstract=abstract, keywords=keywords)
    profile = ty.resolve_profile(template)

    if doc_model.tables:
        pb.write_table_files(ws, doc_model.tables)

    d = build_docx(doc_model, profile)
    out = paper_dir / out_name
    d.save(str(out))

    return {
        "ok": out.is_file(),
        "docx": str(out),
        "bytes": out.stat().st_size if out.is_file() else 0,
        "profile": profile.name,
        "sections": len(doc_model.sections),
        "tables": len(doc_model.tables),
        "figures": len(doc_model.figures),
        "references": len(doc_model.references),
        "budget": doc_model.budget().as_dict(),
        "notes": notes,
    }


__all__ = [
    "build_docx", "export_docx", "set_run_font",
    "set_first_line_indent_chars", "set_cell_borders",
]
