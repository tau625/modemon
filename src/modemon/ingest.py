"""赛题文件解析：把 .txt/.md/.pdf/.docx 统一读成纯文本。

真实使用场景里，赛题往往是官方发的 PDF 或 Word；直接要求用户手工转文本太不友好。
"""

from __future__ import annotations

import re
from pathlib import Path

SUPPORTED = (".txt", ".md", ".markdown", ".pdf", ".docx", ".doc")


class IngestError(RuntimeError):
    pass


def _from_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise IngestError("解析 PDF 需要 pypdf：uv pip install pypdf") from exc

    try:
        reader = PdfReader(str(path))
        pages = [(page.extract_text() or "") for page in reader.pages]
    except Exception as exc:  # noqa: BLE001
        raise IngestError(f"读取 PDF 失败：{exc}") from exc

    text = "\n\n".join(pages).strip()
    if not text:
        raise IngestError("该 PDF 没有可提取的文本层（可能是扫描件，需要 OCR）")
    return text


def _from_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise IngestError("解析 .docx 需要 python-docx：uv pip install python-docx") from exc

    try:
        document = docx.Document(str(path))
    except Exception as exc:  # noqa: BLE001
        raise IngestError(f"读取 Word 失败：{exc}") from exc

    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise IngestError("该 Word 文档没有可提取的文本")
    return text


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{4,}", "\n\n\n", text)
    return text.strip()


def load_problem_text(path: Path) -> str:
    """读取赛题文件为纯文本；不支持的格式给出明确提示。"""
    path = Path(path)
    if not path.exists():
        raise IngestError(f"赛题文件不存在：{path}")
    suffix = path.suffix.lower()

    if suffix in (".txt", ".md", ".markdown"):
        text = path.read_text(encoding="utf-8", errors="replace")
    elif suffix == ".pdf":
        text = _from_pdf(path)
    elif suffix == ".docx":
        text = _from_docx(path)
    elif suffix == ".doc":
        raise IngestError("不支持旧版 .doc，请先用 Word 另存为 .docx 或导出 PDF")
    else:
        raise IngestError(f"不支持的赛题格式 {suffix or '（无扩展名）'}；支持：{', '.join(SUPPORTED)}")

    text = _clean(text)
    if not text:
        raise IngestError(f"赛题文件内容为空：{path}")
    return text
